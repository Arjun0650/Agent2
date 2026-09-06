import json
import hashlib
from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = BASE_DIR / "uploads"

WORKBOOK_FILE = DATA_DIR / "GCGW_Payment_Management.xlsx"
PENDING_FILE = DATA_DIR / "pending_ai_queue.json"

DATA_DIR.mkdir(
    parents=True,
    exist_ok=True
)

UPLOAD_DIR.mkdir(
    parents=True,
    exist_ok=True
)


# ============================================================
# BASIC HELPERS
# ============================================================

def clean(value):
    if value is None:
        return ""

    return str(value).strip()


def normalized(value):
    return clean(value).lower()


def amount_value(value):
    try:
        return float(value or 0)

    except (TypeError, ValueError):
        return 0.0


# ============================================================
# WORKBOOK HELPERS
# ============================================================

def load_book():

    if not WORKBOOK_FILE.exists():
        raise FileNotFoundError(
            f"GCGW workbook not found: {WORKBOOK_FILE}"
        )

    return load_workbook(
        WORKBOOK_FILE
    )


def save_book(workbook):

    workbook.save(
        WORKBOOK_FILE
    )


# ============================================================
# PAYMENT FINGERPRINT
# ============================================================

def payment_fingerprint(payment):

    raw = "|".join([
        normalized(
            payment.get("transaction_id")
        ),

        str(
            round(
                amount_value(
                    payment.get("amount")
                ),
                2
            )
        ),

        normalized(
            payment.get("payment_date")
        ),

        normalized(
            payment.get("paid_to")
        )
    ])

    return hashlib.sha256(
        raw.encode("utf-8")
    ).hexdigest()


# ============================================================
# DUPLICATE CHECK
# ============================================================

def check_duplicate(payment):

    workbook = load_book()

    transaction_id = normalized(
        payment.get("transaction_id")
    )

    amount = round(
        amount_value(
            payment.get("amount")
        ),
        2
    )

    payment_date = normalized(
        payment.get("payment_date")
    )

    paid_to = normalized(
        payment.get("paid_to")
    )

    # ========================================================
    # CHECK ALL PAYMENTS
    # ========================================================

    approved = workbook[
        "All Payments"
    ]

    for row in range(
        2,
        approved.max_row + 1
    ):

        internal_id = approved.cell(
            row,
            1
        ).value

        if not internal_id:
            continue

        # ----------------------------------------------------
        # Reference / transaction ID
        # ----------------------------------------------------

        existing_reference = normalized(
            approved.cell(
                row,
                11
            ).value
        )

        if (
            transaction_id
            and existing_reference
            and transaction_id == existing_reference
        ):

            return {
                "duplicate": True,
                "location": "All Payments",
                "row": row,
                "reason":
                    "Transaction/reference ID already exists "
                    "in approved payments."
            }

        # ----------------------------------------------------
        # Amount + date + payee
        # ----------------------------------------------------

        existing_amount = round(
            amount_value(
                approved.cell(
                    row,
                    4
                ).value
            ),
            2
        )

        existing_date = normalized(
            approved.cell(
                row,
                2
            ).value
        )

        existing_payee = normalized(
            approved.cell(
                row,
                5
            ).value
        )

        if (
            amount > 0
            and payment_date
            and paid_to
            and amount == existing_amount
            and payment_date == existing_date
            and paid_to == existing_payee
        ):

            return {
                "duplicate": True,
                "location": "All Payments",
                "row": row,
                "reason":
                    "Same amount, date and payee already exist "
                    "in approved payments."
            }

    # ========================================================
    # CHECK NEEDS REVIEW
    # ========================================================

    review = workbook[
        "Needs Review"
    ]

    for row in range(
        2,
        review.max_row + 1
    ):

        review_id = normalized(
            review.cell(
                row,
                1
            ).value
        )

        if not review_id:
            continue

        decision = normalized(
            review.cell(
                row,
                8
            ).value
        )

        # Important:
        # Ignore completed rows and the row currently
        # being approved.

        if decision in {
            "approved",
            "rejected",
            "deleted",
            "approval in progress"
        }:
            continue

        # ----------------------------------------------------
        # Transaction/reference ID
        # ----------------------------------------------------

        if (
            transaction_id
            and review_id
            and not review_id.startswith("pending-")
            and transaction_id == review_id
        ):

            return {
                "duplicate": True,
                "location": "Needs Review",
                "row": row,
                "reason":
                    "Transaction already exists in Needs Review."
            }

        # ----------------------------------------------------
        # Amount + date + payee
        # ----------------------------------------------------

        review_amount = round(
            amount_value(
                review.cell(
                    row,
                    3
                ).value
            ),
            2
        )

        review_date = normalized(
            review.cell(
                row,
                2
            ).value
        )

        review_payee = normalized(
            review.cell(
                row,
                4
            ).value
        )

        if (
            amount > 0
            and payment_date
            and paid_to
            and amount == review_amount
            and payment_date == review_date
            and paid_to == review_payee
        ):

            return {
                "duplicate": True,
                "location": "Needs Review",
                "row": row,
                "reason":
                    "Same amount, date and payee are already "
                    "waiting for review."
            }

    return {
        "duplicate": False,
        "location": None,
        "row": None,
        "reason": None
    }


# ============================================================
# ADD APPROVED PAYMENT
# ============================================================

def add_approved_payment(
    payment,
    screenshot_filename
):

    duplicate = check_duplicate(
        payment
    )

    if duplicate["duplicate"]:

        return {
            "success": False,
            "status": "DUPLICATE",
            "duplicate": duplicate
        }

    workbook = load_book()

    sheet = workbook[
        "All Payments"
    ]

    # Find first empty row

    row = 2

    while sheet.cell(
        row,
        1
    ).value:

        row += 1

    now = datetime.now()

    internal_id = (
        "GCGW-"
        + now.strftime(
            "%Y%m%d%H%M%S%f"
        )
    )

    fingerprint = payment_fingerprint(
        payment
    )

    values = [
        internal_id,
        payment.get("payment_date"),
        payment.get("payment_time"),
        amount_value(
            payment.get("amount")
        ),
        payment.get("paid_to"),
        payment.get("project"),
        payment.get("category"),
        payment.get("purpose"),
        payment.get("payment_mode"),
        payment.get("payment_app"),
        payment.get("transaction_id"),
        screenshot_filename,
        payment.get("confidence"),
        "Approved",
        fingerprint
    ]

    for column, value in enumerate(
        values,
        start=1
    ):

        sheet.cell(
            row,
            column
        ).value = value

    save_book(
        workbook
    )

    return {
        "success": True,
        "status": "APPROVED",
        "row": row,
        "internal_id": internal_id
    }


# ============================================================
# ADD PAYMENT TO NEEDS REVIEW
# ============================================================

def add_to_review(
    payment,
    screenshot_filename,
    problems=None
):

    duplicate = check_duplicate(
        payment
    )

    if duplicate["duplicate"]:

        return {
            "success": False,
            "status": "DUPLICATE",
            "duplicate": duplicate
        }

    workbook = load_book()

    sheet = workbook[
        "Needs Review"
    ]

    row = 2

    while sheet.cell(
        row,
        1
    ).value:

        row += 1

    transaction_id = clean(
        payment.get(
            "transaction_id"
        )
    )

    # --------------------------------------------------------
    # Generate temporary ID if Gemini could not read one
    # --------------------------------------------------------

    if not transaction_id:

        fingerprint = payment_fingerprint(
            payment
        )[:12]

        transaction_id = (
            "PENDING-"
            + fingerprint
        )

    # --------------------------------------------------------
    # Review reasons
    # --------------------------------------------------------

    reasons = list(
        problems or []
    )

    ai_reason = clean(
        payment.get(
            "review_reason"
        )
    )

    if (
        ai_reason
        and ai_reason not in reasons
    ):

        reasons.append(
            ai_reason
        )

    reason_text = (
        "; ".join(reasons)
        if reasons
        else "Manual review required"
    )

    # ========================================================
    # ORIGINAL REVIEW COLUMNS
    # ========================================================

    # 1 Transaction ID
    sheet.cell(
        row,
        1
    ).value = transaction_id

    # 2 Date
    sheet.cell(
        row,
        2
    ).value = payment.get(
        "payment_date"
    )

    # 3 Amount
    sheet.cell(
        row,
        3
    ).value = amount_value(
        payment.get("amount")
    )

    # 4 Paid To
    sheet.cell(
        row,
        4
    ).value = payment.get(
        "paid_to"
    )

    # 5 Project
    sheet.cell(
        row,
        5
    ).value = payment.get(
        "project"
    )

    # 6 Category
    sheet.cell(
        row,
        6
    ).value = payment.get(
        "category"
    )

    # 7 Review reason
    sheet.cell(
        row,
        7
    ).value = reason_text

    # 8 Final decision
    sheet.cell(
        row,
        8
    ).value = "Pending"

    # ========================================================
    # EXTRA AUDIT COLUMNS
    # ========================================================

    # 9 Screenshot File
    sheet.cell(
        row,
        9
    ).value = screenshot_filename

    # 10 Time
    sheet.cell(
        row,
        10
    ).value = payment.get(
        "payment_time"
    )

    # 11 Payment Mode
    sheet.cell(
        row,
        11
    ).value = payment.get(
        "payment_mode"
    )

    # 12 Payment App
    sheet.cell(
        row,
        12
    ).value = payment.get(
        "payment_app"
    )

    # 13 Purpose
    sheet.cell(
        row,
        13
    ).value = payment.get(
        "purpose"
    )

    # 14 Confidence
    sheet.cell(
        row,
        14
    ).value = payment.get(
        "confidence"
    )

    save_book(
        workbook
    )

    return {
        "success": True,
        "status": "NEEDS_REVIEW",
        "row": row,
        "transaction_id": transaction_id
    }


# ============================================================
# PENDING AI QUEUE
# ============================================================

def load_pending_queue():

    if not PENDING_FILE.exists():
        return []

    try:

        data = json.loads(
            PENDING_FILE.read_text(
                encoding="utf-8"
            )
        )

        if isinstance(
            data,
            list
        ):
            return data

        return []

    except Exception:
        return []


def save_pending_queue(queue):

    PENDING_FILE.write_text(
        json.dumps(
            queue,
            indent=2,
            ensure_ascii=False
        ),
        encoding="utf-8"
    )


def add_pending_ai(
    screenshot_filename,
    error
):

    queue = load_pending_queue()

    # --------------------------------------------------------
    # Do not add same screenshot repeatedly
    # --------------------------------------------------------

    for item in queue:

        if (
            item.get("filename")
            == screenshot_filename
            and normalized(
                item.get("status")
            ) == "pending"
        ):

            # Update latest error information

            item["error"] = clean(
                error
            )

            item["last_attempt_at"] = (
                datetime.now().isoformat(
                    timespec="seconds"
                )
            )

            save_pending_queue(
                queue
            )

            return {
                "success": True,
                "status": "PENDING_AI",
                "already_exists": True
            }

    queue.append({
        "filename":
            screenshot_filename,

        "status":
            "Pending",

        "error":
            clean(error),

        "added_at":
            datetime.now().isoformat(
                timespec="seconds"
            )
    })

    save_pending_queue(
        queue
    )

    return {
        "success": True,
        "status": "PENDING_AI",
        "already_exists": False
    }


# ============================================================
# MAIN PAYMENT DECISION ENGINE
# ============================================================

def process_extracted_payment(
    payment,
    validation,
    screenshot_filename
):

    # Make sure payment is always a dictionary.

    if payment is None:
        payment = {}

    # --------------------------------------------------------
    # Duplicate check
    # --------------------------------------------------------

    duplicate = check_duplicate(
        payment
    )

    if duplicate["duplicate"]:

        return {
            "success": True,
            "status": "DUPLICATE",
            "duplicate": duplicate
        }

    # --------------------------------------------------------
    # Validation problems
    # --------------------------------------------------------

    if validation:

        problems = list(
            validation.get(
                "problems",
                []
            )
            or []
        )

    else:

        problems = [
            "Payment validation unavailable"
        ]

    # --------------------------------------------------------
    # Determine whether manual review is required
    # --------------------------------------------------------

    needs_review = True

    if validation:

        needs_review = validation.get(
            "needs_review",
            True
        )

    # --------------------------------------------------------
    # PARTIAL GEMINI EXTRACTION
    #
    # This is important:
    #
    # Missing project/category/purpose does NOT mean Pending AI.
    # It means Needs Review.
    # --------------------------------------------------------

    if needs_review:

        return add_to_review(
            payment,
            screenshot_filename,
            problems
        )

    # --------------------------------------------------------
    # Fully validated payment
    # --------------------------------------------------------

    return add_approved_payment(
        payment,
        screenshot_filename
    )


# ============================================================
# APPROVE NEEDS REVIEW PAYMENT
# ============================================================

def approve_review_payment(
    review_row,
    corrections
):

    workbook = load_book()

    review_sheet = workbook[
        "Needs Review"
    ]

    # ========================================================
    # VALIDATE ROW
    # ========================================================

    if (
        review_row < 2
        or review_row > review_sheet.max_row
    ):

        return {
            "success": False,
            "status": "FAILED",
            "error": "Invalid review row."
        }

    current_decision = normalized(
        review_sheet.cell(
            review_row,
            8
        ).value
    )

    if current_decision in {
        "approved",
        "rejected",
        "deleted"
    }:

        return {
            "success": False,
            "status": "FAILED",
            "error":
                "This review item has already been processed."
        }

    # ========================================================
    # EXISTING REVIEW DATA
    # ========================================================

    old_transaction_id = clean(
        review_sheet.cell(
            review_row,
            1
        ).value
    )

    screenshot_filename = clean(
        review_sheet.cell(
            review_row,
            9
        ).value
    )

    # ========================================================
    # TRANSACTION ID
    # ========================================================

    transaction_id = clean(
        corrections.get(
            "transaction_id"
        )
    )

    if not transaction_id:

        transaction_id = old_transaction_id

    # Generated PENDING IDs cannot be used as real
    # transaction/reference IDs.

    if normalized(
        transaction_id
    ).startswith(
        "pending-"
    ):

        transaction_id = ""

    # ========================================================
    # DATE COMPATIBILITY
    #
    # Supports both:
    # payment_date
    # date
    # ========================================================

    corrected_date = corrections.get(
        "payment_date"
    )

    if corrected_date is None:

        corrected_date = corrections.get(
            "date"
        )

    if corrected_date is None:

        corrected_date = review_sheet.cell(
            review_row,
            2
        ).value

    # ========================================================
    # BUILD CORRECTED PAYMENT
    # ========================================================

    payment = {

        "amount":
            amount_value(
                corrections.get(
                    "amount",
                    review_sheet.cell(
                        review_row,
                        3
                    ).value
                )
            ),

        "paid_to":
            clean(
                corrections.get(
                    "paid_to",
                    review_sheet.cell(
                        review_row,
                        4
                    ).value
                )
            ),

        "payment_date":
            clean(
                corrected_date
            ),

        "payment_time":
            clean(
                corrections.get(
                    "payment_time",
                    review_sheet.cell(
                        review_row,
                        10
                    ).value
                )
            ),

        "transaction_id":
            transaction_id,

        "payment_mode":
            clean(
                corrections.get(
                    "payment_mode",
                    review_sheet.cell(
                        review_row,
                        11
                    ).value
                )
            ),

        "payment_app":
            clean(
                corrections.get(
                    "payment_app",
                    review_sheet.cell(
                        review_row,
                        12
                    ).value
                )
            ),

        "purpose":
            clean(
                corrections.get(
                    "purpose",
                    review_sheet.cell(
                        review_row,
                        13
                    ).value
                )
            ),

        "project":
            clean(
                corrections.get(
                    "project",
                    review_sheet.cell(
                        review_row,
                        5
                    ).value
                )
            ),

        "category":
            clean(
                corrections.get(
                    "category",
                    review_sheet.cell(
                        review_row,
                        6
                    ).value
                )
            ),

        "confidence":
            amount_value(
                review_sheet.cell(
                    review_row,
                    14
                ).value
            ),

        "needs_review":
            False,

        "review_reason":
            None
    }

    # ========================================================
    # STRICT APPROVAL VALIDATION
    # ========================================================

    missing = []

    if payment["amount"] <= 0:
        missing.append(
            "amount"
        )

    if not payment["paid_to"]:
        missing.append(
            "paid_to"
        )

    if not payment["payment_date"]:
        missing.append(
            "payment_date"
        )

    if not payment["transaction_id"]:
        missing.append(
            "transaction/reference ID"
        )

    if not payment["project"]:
        missing.append(
            "project"
        )

    if not payment["category"]:
        missing.append(
            "category"
        )

    if missing:

        return {
            "success": False,
            "status": "NEEDS_REVIEW",
            "error":
                "Cannot approve until these fields are confirmed: "
                + ", ".join(missing)
        }

    # ========================================================
    # TEMPORARILY MARK THIS REVIEW ROW
    #
    # Prevents duplicate detection from matching the row
    # against itself.
    # ========================================================

    original_decision = review_sheet.cell(
        review_row,
        8
    ).value

    review_sheet.cell(
        review_row,
        8
    ).value = "Approval In Progress"

    save_book(
        workbook
    )

    # ========================================================
    # DUPLICATE CHECK
    # ========================================================

    try:

        duplicate = check_duplicate(
            payment
        )

    except Exception as error:

        workbook = load_book()

        review_sheet = workbook[
            "Needs Review"
        ]

        review_sheet.cell(
            review_row,
            8
        ).value = original_decision

        save_book(
            workbook
        )

        return {
            "success": False,
            "status": "FAILED",
            "error":
                "Duplicate check failed: "
                + str(error)
        }

    # ========================================================
    # REAL DUPLICATE FOUND
    # ========================================================

    if duplicate["duplicate"]:

        workbook = load_book()

        review_sheet = workbook[
            "Needs Review"
        ]

        review_sheet.cell(
            review_row,
            8
        ).value = original_decision

        save_book(
            workbook
        )

        return {
            "success": False,
            "status": "DUPLICATE",
            "duplicate": duplicate,
            "error":
                duplicate.get(
                    "reason"
                )
        }

    # ========================================================
    # WRITE APPROVED PAYMENT
    # ========================================================

    workbook = load_book()

    approved_sheet = workbook[
        "All Payments"
    ]

    review_sheet = workbook[
        "Needs Review"
    ]

    target_row = 2

    while approved_sheet.cell(
        target_row,
        1
    ).value:

        target_row += 1

    now = datetime.now()

    internal_id = (
        "GCGW-"
        + now.strftime(
            "%Y%m%d%H%M%S%f"
        )
    )

    values = [
        internal_id,
        payment["payment_date"],
        payment["payment_time"],
        payment["amount"],
        payment["paid_to"],
        payment["project"],
        payment["category"],
        payment["purpose"],
        payment["payment_mode"],
        payment["payment_app"],
        payment["transaction_id"],
        screenshot_filename,
        payment["confidence"],
        "Approved",
        payment_fingerprint(
            payment
        )
    ]

    for column, value in enumerate(
        values,
        start=1
    ):

        approved_sheet.cell(
            target_row,
            column
        ).value = value

    # ========================================================
    # MARK REVIEW RECORD APPROVED
    # ========================================================

    review_sheet.cell(
        review_row,
        8
    ).value = "Approved"

    # ========================================================
    # SAVE CORRECTED REVIEW VALUES
    # ========================================================

    review_sheet.cell(
        review_row,
        1
    ).value = payment[
        "transaction_id"
    ]

    review_sheet.cell(
        review_row,
        2
    ).value = payment[
        "payment_date"
    ]

    review_sheet.cell(
        review_row,
        3
    ).value = payment[
        "amount"
    ]

    review_sheet.cell(
        review_row,
        4
    ).value = payment[
        "paid_to"
    ]

    review_sheet.cell(
        review_row,
        5
    ).value = payment[
        "project"
    ]

    review_sheet.cell(
        review_row,
        6
    ).value = payment[
        "category"
    ]

    review_sheet.cell(
        review_row,
        10
    ).value = payment[
        "payment_time"
    ]

    review_sheet.cell(
        review_row,
        11
    ).value = payment[
        "payment_mode"
    ]

    review_sheet.cell(
        review_row,
        12
    ).value = payment[
        "payment_app"
    ]

    review_sheet.cell(
        review_row,
        13
    ).value = payment[
        "purpose"
    ]

    # ========================================================
    # SAVE WORKBOOK
    # ========================================================

    save_book(
        workbook
    )

    return {
        "success": True,
        "status": "APPROVED",
        "internal_id": internal_id,
        "approved_row": target_row,
        "review_row": review_row
    }
