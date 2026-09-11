import os
import json
from pathlib import Path
from datetime import date, datetime

from fastapi import FastAPI, UploadFile, File, Request
from fastapi.responses import (
    HTMLResponse,
    JSONResponse,
    FileResponse,
    RedirectResponse
)
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware

from openpyxl import load_workbook

from ai_engine import extract_payment_for_web
from payment_engine import (
    process_extracted_payment,
    add_pending_ai,
    approve_review_payment
)

from auth import (
    register_user,
    authenticate_user,
    get_user_by_id
)


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

TEMPLATES_DIR = BASE_DIR / "templates"
STATIC_DIR = BASE_DIR / "static"
UPLOAD_DIR = BASE_DIR / "uploads"
DATA_DIR = BASE_DIR / "data"

WORKBOOK_FILE = DATA_DIR / "GCGW_Payment_Management.xlsx"
PENDING_FILE = DATA_DIR / "pending_ai_queue.json"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
DATA_DIR.mkdir(parents=True, exist_ok=True)
STATIC_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# FASTAPI
# ============================================================

app = FastAPI(
    title="GCGW Payment Management Agent",
    description=(
        "AI payment and expenditure management "
        "for Gopal Chavan Guniting Work."
    ),
    version="1.3.0"
)


# ============================================================
# LOGIN SESSION
# ============================================================

SESSION_SECRET = os.environ.get(
    "SESSION_SECRET_KEY",
    "gcgw-development-secret-change-on-render"
)

app.mount(
    "/static",
    StaticFiles(directory=str(STATIC_DIR)),
    name="static"
)


# ============================================================
# BASIC HELPERS
# ============================================================

def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


def safe_amount(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def normalize_date(value):

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if not value:
        return None

    text = str(value).strip()

    for fmt in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%d %b %Y",
        "%d %B %Y"
    ):
        try:
            return datetime.strptime(
                text,
                fmt
            ).date()

        except ValueError:
            pass

    return None


def load_gcgw_workbook():

    if not WORKBOOK_FILE.exists():

        raise FileNotFoundError(
            f"GCGW workbook is unavailable: {WORKBOOK_FILE}"
        )

    return load_workbook(
        WORKBOOK_FILE,
        data_only=False
    )


# ============================================================
# APPROVED PAYMENTS
# ============================================================

def get_approved_payments():

    workbook = load_gcgw_workbook()

    sheet = workbook["All Payments"]

    records = []

    for row in range(
        2,
        sheet.max_row + 1
    ):

        internal_id = sheet.cell(
            row,
            1
        ).value

        if not internal_id:
            continue

        status = clean_text(
            sheet.cell(
                row,
                14
            ).value
        )

        if status.lower() != "approved":
            continue

        payment_date = normalize_date(
            sheet.cell(
                row,
                2
            ).value
        )

        records.append({

            "id":
                clean_text(
                    internal_id
                ),

            "date":
                (
                    payment_date.isoformat()
                    if payment_date
                    else clean_text(
                        sheet.cell(
                            row,
                            2
                        ).value
                    )
                ),

            "time":
                clean_text(
                    sheet.cell(
                        row,
                        3
                    ).value
                ),

            "amount":
                safe_amount(
                    sheet.cell(
                        row,
                        4
                    ).value
                ),

            "paid_to":
                clean_text(
                    sheet.cell(
                        row,
                        5
                    ).value
                ),

            "project":
                clean_text(
                    sheet.cell(
                        row,
                        6
                    ).value
                ),

            "category":
                clean_text(
                    sheet.cell(
                        row,
                        7
                    ).value
                ),

            "purpose":
                clean_text(
                    sheet.cell(
                        row,
                        8
                    ).value
                ),

            "payment_mode":
                clean_text(
                    sheet.cell(
                        row,
                        9
                    ).value
                ),

            "payment_app":
                clean_text(
                    sheet.cell(
                        row,
                        10
                    ).value
                ),

            "reference_no":
                clean_text(
                    sheet.cell(
                        row,
                        11
                    ).value
                ),

            "screenshot":
                clean_text(
                    sheet.cell(
                        row,
                        12
                    ).value
                ),

            "confidence":
                safe_amount(
                    sheet.cell(
                        row,
                        13
                    ).value
                ),

            "status":
                status
        })

    workbook.close()

    return records


# ============================================================
# NEEDS REVIEW
# ============================================================

def get_review_payments():

    workbook = load_gcgw_workbook()

    sheet = workbook["Needs Review"]

    records = []

    for row in range(
        2,
        sheet.max_row + 1
    ):

        transaction_id = sheet.cell(
            row,
            1
        ).value

        if not transaction_id:
            continue

        decision = clean_text(
            sheet.cell(
                row,
                8
            ).value
        )

        if decision.lower() in {
            "approved",
            "rejected",
            "deleted"
        }:
            continue

        records.append({

            "row":
                row,

            "transaction_id":
                clean_text(
                    transaction_id
                ),

            "date":
                clean_text(
                    sheet.cell(
                        row,
                        2
                    ).value
                ),

            "amount":
                safe_amount(
                    sheet.cell(
                        row,
                        3
                    ).value
                ),

            "paid_to":
                clean_text(
                    sheet.cell(
                        row,
                        4
                    ).value
                ),

            "project":
                clean_text(
                    sheet.cell(
                        row,
                        5
                    ).value
                ),

            "category":
                clean_text(
                    sheet.cell(
                        row,
                        6
                    ).value
                ),

            "reason":
                clean_text(
                    sheet.cell(
                        row,
                        7
                    ).value
                ),

            "decision":
                decision or "Pending",

            "screenshot":
                clean_text(
                    sheet.cell(
                        row,
                        9
                    ).value
                ),

            "time":
                clean_text(
                    sheet.cell(
                        row,
                        10
                    ).value
                ),

            "payment_mode":
                clean_text(
                    sheet.cell(
                        row,
                        11
                    ).value
                ),

            "payment_app":
                clean_text(
                    sheet.cell(
                        row,
                        12
                    ).value
                ),

            "purpose":
                clean_text(
                    sheet.cell(
                        row,
                        13
                    ).value
                ),

            "confidence":
                safe_amount(
                    sheet.cell(
                        row,
                        14
                    ).value
                )
        })

    workbook.close()

    return records


# ============================================================
# PENDING AI
# ============================================================

def get_pending_ai():

    if not PENDING_FILE.exists():
        return []

    try:

        data = json.loads(
            PENDING_FILE.read_text(
                encoding="utf-8"
            )
        )

        if isinstance(data, list):
            return data

    except Exception:
        pass

    return []


def save_pending_ai(records):

    try:

        PENDING_FILE.write_text(
            json.dumps(
                records,
                indent=2,
                ensure_ascii=False,
                default=str
            ),
            encoding="utf-8"
        )

    except Exception:
        pass


def remove_pending_item(filename):

    pending = get_pending_ai()

    updated = []

    for item in pending:

        item_filename = clean_text(
            item.get("filename")
        )

        if item_filename != filename:
            updated.append(item)

    save_pending_ai(updated)


# ============================================================
# ACCOUNTING SUMMARY
# ============================================================

def accounting_summary():

    payments = get_approved_payments()
    review = get_review_payments()
    pending = get_pending_ai()

    today = date.today()

    total = 0.0
    today_total = 0.0
    month_total = 0.0

    project_totals = {}
    category_totals = {}

    for payment in payments:

        amount = safe_amount(
            payment.get("amount")
        )

        total += amount

        payment_date = normalize_date(
            payment.get("date")
        )

        if payment_date:

            if payment_date == today:
                today_total += amount

            if (
                payment_date.year == today.year
                and
                payment_date.month == today.month
            ):
                month_total += amount

        project = (
            payment.get("project")
            or "Unassigned"
        )

        category = (
            payment.get("category")
            or "Other"
        )

        project_totals[project] = (
            project_totals.get(
                project,
                0
            )
            + amount
        )

        category_totals[category] = (
            category_totals.get(
                category,
                0
            )
            + amount
        )

    return {

        "total_expenditure":
            round(
                total,
                2
            ),

        "this_month":
            round(
                month_total,
                2
            ),

        "today":
            round(
                today_total,
                2
            ),

        "approved_payments":
            len(payments),

        "needs_review":
            len(review),

        "pending_ai":
            len(pending),

        "project_totals":
            project_totals,

        "category_totals":
            category_totals
    }


# ============================================================
# AUTH PAGES
# ============================================================

@app.get(
    "/",
    response_class=HTMLResponse
)
def login_page(request: Request):

    if request.session.get("user_id"):

        return RedirectResponse(
            url="/dashboard",
            status_code=302
        )

    login_file = (
        TEMPLATES_DIR /
        "login.html"
    )

    return login_file.read_text(
        encoding="utf-8"
    )


@app.get(
    "/register",
    response_class=HTMLResponse
)
def register_page(request: Request):

    if request.session.get("user_id"):

        return RedirectResponse(
            url="/dashboard",
            status_code=302
        )

    register_file = (
        TEMPLATES_DIR /
        "register.html"
    )

    return register_file.read_text(
        encoding="utf-8"
    )


@app.get(
    "/dashboard",
    response_class=HTMLResponse
)
def dashboard(request: Request):

    if not request.session.get("user_id"):

        return RedirectResponse(
            url="/",
            status_code=302
        )

    index_file = (
        TEMPLATES_DIR /
        "index.html"
    )

    return index_file.read_text(
        encoding="utf-8"
    )


# ============================================================
# REGISTER API
# ============================================================

@app.post("/api/register")
async def register_account(
    request: Request
):

    try:

        payload = await request.json()

        full_name = (
            payload.get("full_name")
            or ""
        )

        email = (
            payload.get("email")
            or ""
        )

        password = (
            payload.get("password")
            or ""
        )

        result = register_user(
            full_name,
            email,
            password
        )

        if not result.get("success"):

            return JSONResponse(
                status_code=400,
                content=result
            )

        user = result["user"]

        request.session["user_id"] = (
            user["id"]
        )

        return {
            "success": True,
            "message": (
                "Account created successfully."
            ),
            "user": user
        }

    except Exception as e:

        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": str(e)
            }
        )


# ============================================================
# LOGIN API
# ============================================================

@app.post("/api/login")
async def login_account(
    request: Request
):

    try:

        payload = await request.json()

        email = (
            payload.get("email")
            or ""
        )

        password = (
            payload.get("password")
            or ""
        )

        result = authenticate_user(
            email,
            password
        )

        if not result.get("success"):

            return JSONResponse(
                status_code=401,
                content=result
            )

        user = result["user"]

        request.session["user_id"] = (
            user["id"]
        )

        return {
            "success": True,
            "message": "Login successful.",
            "user": user
        }

    except Exception as e:

        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": str(e)
            }
        )


# ============================================================
# LOGOUT
# ============================================================

@app.post("/api/logout")
async def logout_account(
    request: Request
):

    request.session.clear()

    return {
        "success": True,
        "message": "Logged out successfully."
    }


# ============================================================
# CURRENT USER
# ============================================================

@app.get("/api/me")
def current_user(
    request: Request
):

    user_id = (
        request.session.get(
            "user_id"
        )
    )

    if not user_id:

        return JSONResponse(
            status_code=401,
            content={
                "success": False,
                "error": "Not authenticated."
            }
        )

    user = get_user_by_id(
        user_id
    )

    if not user:

        request.session.clear()

        return JSONResponse(
            status_code=401,
            content={
                "success": False,
                "error": "User account not found."
            }
        )

    return {
        "success": True,
        "user": user
    }


# ============================================================
# PROTECT DASHBOARD + PAYMENT APIs
# ============================================================

@app.middleware("http")
async def authentication_middleware(
    request: Request,
    call_next
):

    path = request.url.path

    public_paths = {
        "/",
        "/register",
        "/api/login",
        "/api/register",
        "/health"
    }

    if (
        path in public_paths
        or path.startswith("/static/")
    ):

        return await call_next(
            request
        )

    protected = (
        path == "/dashboard"
        or path.startswith("/api/")
    )

    if (
        protected
        and not request.session.get("user_id")
    ):

        if path.startswith("/api/"):

            return JSONResponse(
                status_code=401,
                content={
                    "success": False,
                    "error": (
                        "Please login to continue."
                    )
                }
            )

        return RedirectResponse(
            url="/",
            status_code=302
        )

    return await call_next(
        request
    )


# SessionMiddleware must wrap the authentication middleware
# so request.session is available during auth checks.
app.add_middleware(
    SessionMiddleware,
    secret_key=SESSION_SECRET,
    same_site="lax",
    https_only=True
)


# ============================================================
# HEALTH
# ============================================================

@app.get("/health")
def health():

    workbook_ready = (
        WORKBOOK_FILE.exists()
    )

    api_ready = bool(
        os.environ.get(
            "GOOGLE_API_KEY"
        )
    )

    return {

        "status":
            (
                "healthy"
                if workbook_ready
                else "degraded"
            ),

        "application":
            "GCGW Payment Management Agent",

        "workbook":
            workbook_ready,

        "gemini_api_configured":
            api_ready,

        "version":
            "1.3.0"
    }


# ============================================================
# SUMMARY
# ============================================================

@app.get("/api/summary")
def summary():

    try:

        return accounting_summary()

    except Exception as e:

        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": str(e)
            }
        )


# ============================================================
# PAYMENT HISTORY
# ============================================================

@app.get("/api/payments")
def payment_history():

    try:

        payments = (
            get_approved_payments()
        )

        payments.reverse()

        return {
            "success": True,
            "count": len(payments),
            "payments": payments
        }

    except Exception as e:

        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": str(e)
            }
        )


# ============================================================
# REVIEW QUEUE
# ============================================================

@app.get("/api/review")
def review_queue():

    try:

        payments = (
            get_review_payments()
        )

        return {
            "success": True,
            "count": len(payments),
            "payments": payments
        }

    except Exception as e:

        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": str(e)
            }
        )


# ============================================================
# APPROVE REVIEW PAYMENT
# ============================================================

@app.post(
    "/api/review/{review_row}/approve"
)
async def approve_review(
    review_row: int,
    payload: dict
):

    try:

        # Compatibility with frontend.
        # Frontend may send "date".
        if (
            "date" in payload
            and
            "payment_date" not in payload
        ):
            payload["payment_date"] = (
                payload.get("date")
            )

        result = (
            approve_review_payment(
                review_row,
                payload
            )
        )

        if not result.get(
            "success"
        ):

            return JSONResponse(
                status_code=400,
                content=result
            )

        return result

    except Exception as e:

        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "status": "FAILED",
                "error": str(e)
            }
        )


# ============================================================
# PENDING AI LIST
# ============================================================

@app.get("/api/pending-ai")
def pending_ai():

    payments = get_pending_ai()

    return {
        "success": True,
        "count": len(payments),
        "payments": payments
    }


# ============================================================
# PROCESS ONE SCREENSHOT
# ============================================================

def process_screenshot_file(
    destination: Path
):

    stored_name = destination.name

    # --------------------------------------------------------
    # ONE FAST AI EXTRACTION
    # --------------------------------------------------------

    try:

        extraction = (
            extract_payment_for_web(
                destination
            )
        )

    except Exception as e:

        extraction = {
            "success": False,
            "payment": None,
            "validation": None,
            "quota_error": False,
            "error": str(e)
        }

    # --------------------------------------------------------
    # COMPLETE AI FAILURE
    # --------------------------------------------------------

    if not extraction.get(
        "success"
    ):

        error_text = (
            extraction.get("error")
            or
            "AI extraction failed."
        )

        pending_result = (
            add_pending_ai(
                stored_name,
                error_text
            )
        )

        return {

            "filename":
                stored_name,

            "status":
                "PENDING_AI",

            "quota_error":
                extraction.get(
                    "quota_error",
                    False
                ),

            "error":
                error_text,

            "pending":
                pending_result
        }

    # --------------------------------------------------------
    # EXTRACTED PAYMENT
    # --------------------------------------------------------

    payment = (
        extraction.get(
            "payment"
        )
        or {}
    )

    validation = (
        extraction.get(
            "validation"
        )
        or {}
    )

    # --------------------------------------------------------
    # IMPORTANT:
    # PARTIAL DATA SHOULD GO TO REVIEW
    # --------------------------------------------------------

    problems = list(
        validation.get(
            "problems",
            []
        )
        or []
    )

    # Project/category/purpose are allowed
    # to be manually completed later.

    if not payment.get("project"):

        if (
            "Project requires confirmation"
            not in problems
        ):
            problems.append(
                "Project requires confirmation"
            )

    if not payment.get("category"):

        if (
            "Category requires confirmation"
            not in problems
        ):
            problems.append(
                "Category requires confirmation"
            )

    if not payment.get("purpose"):

        if (
            "Purpose can be entered manually"
            not in problems
        ):
            problems.append(
                "Purpose can be entered manually"
            )

    # Financially important fields should
    # still be manually checked if missing.

    if not payment.get("paid_to"):

        if (
            "Payee requires confirmation"
            not in problems
        ):
            problems.append(
                "Payee requires confirmation"
            )

    if safe_amount(
        payment.get("amount")
    ) <= 0:

        if (
            "Amount requires confirmation"
            not in problems
        ):
            problems.append(
                "Amount requires confirmation"
            )

    if not payment.get(
        "payment_date"
    ):

        if (
            "Payment date requires confirmation"
            not in problems
        ):
            problems.append(
                "Payment date requires confirmation"
            )

    if not payment.get(
        "transaction_id"
    ):

        if (
            "Transaction/reference ID requires confirmation"
            not in problems
        ):
            problems.append(
                "Transaction/reference ID requires confirmation"
            )

    # --------------------------------------------------------
    # USER REQUEST:
    # EXTRACTED SCREENSHOTS GO TO REVIEW
    # --------------------------------------------------------

    validation = {
        "valid": False,
        "needs_review": True,
        "problems": problems
    }

    try:

        decision = (
            process_extracted_payment(
                payment=payment,
                validation=validation,
                screenshot_filename=
                    stored_name
            )
        )

    except Exception as e:

        return {

            "filename":
                stored_name,

            "status":
                "FAILED",

            "error":
                (
                    "Payment processing failed: "
                    + str(e)
                )
        }

    # Remove old pending entry if
    # screenshot was successfully processed.

    status = clean_text(
        decision.get(
            "status"
        )
    ).upper()

    if status in {
        "APPROVED",
        "NEEDS_REVIEW",
        "DUPLICATE"
    }:

        remove_pending_item(
            stored_name
        )

    return {

        "filename":
            stored_name,

        "status":
            status or "FAILED",

        "payment":
            payment,

        "validation":
            validation,

        "decision":
            decision
    }


# ============================================================
# RETRY ONE PENDING AI SCREENSHOT
# ============================================================

@app.post(
    "/api/pending-ai/retry/{filename}"
)
async def retry_pending_ai(
    filename: str
):

    try:

        safe_name = Path(
            filename
        ).name

        image_path = (
            UPLOAD_DIR /
            safe_name
        )

        if not image_path.exists():

            return JSONResponse(
                status_code=404,
                content={
                    "success": False,
                    "status": "FILE_MISSING",
                    "error": (
                        "The screenshot file is no longer "
                        "available on the server. "
                        "Please upload it again."
                    )
                }
            )

        result = (
            process_screenshot_file(
                image_path
            )
        )

        return {
            "success": True,
            "result": result
        }

    except Exception as e:

        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": str(e)
            }
        )


# ============================================================
# RETRY ALL PENDING AI SCREENSHOTS
# ============================================================

@app.post(
    "/api/pending-ai/retry-all"
)
async def retry_all_pending_ai():

    pending = get_pending_ai()

    results = []

    processed = 0
    still_pending = 0
    missing = 0
    failed = 0

    # Prevent duplicate filename retries.
    seen = set()

    for item in pending:

        filename = clean_text(
            item.get("filename")
        )

        if not filename:
            continue

        if filename in seen:
            continue

        seen.add(filename)

        image_path = (
            UPLOAD_DIR /
            Path(filename).name
        )

        if not image_path.exists():

            results.append({
                "filename":
                    filename,

                "status":
                    "FILE_MISSING",

                "error":
                    "Screenshot must be uploaded again."
            })

            missing += 1
            continue

        try:

            result = (
                process_screenshot_file(
                    image_path
                )
            )

            results.append(
                result
            )

            status = clean_text(
                result.get("status")
            ).upper()

            if status in {
                "APPROVED",
                "NEEDS_REVIEW",
                "DUPLICATE"
            }:

                processed += 1

            elif status == "PENDING_AI":

                still_pending += 1

            else:

                failed += 1

        except Exception as e:

            results.append({
                "filename":
                    filename,

                "status":
                    "FAILED",

                "error":
                    str(e)
            })

            failed += 1

    return {

        "success": True,

        "processed":
            processed,

        "still_pending":
            still_pending,

        "missing":
            missing,

        "failed":
            failed,

        "results":
            results
    }


# ============================================================
# DOWNLOAD UPDATED EXCEL
# ============================================================

@app.get(
    "/api/download-excel"
)
def download_excel():

    try:

        if not WORKBOOK_FILE.exists():

            return JSONResponse(
                status_code=404,
                content={
                    "success": False,
                    "error":
                        "GCGW Excel workbook was not found."
                }
            )

        timestamp = (
            datetime.now().strftime(
                "%Y-%m-%d_%H-%M-%S"
            )
        )

        download_name = (
            "GCGW_Payment_Management_"
            f"{timestamp}.xlsx"
        )

        return FileResponse(
            path=str(
                WORKBOOK_FILE
            ),
            filename=download_name,
            media_type=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    except Exception as e:

        return JSONResponse(
            status_code=500,
            content={
                "success": False,
                "error": str(e)
            }
        )


# ============================================================
# SCREENSHOT UPLOAD
# ============================================================

@app.post("/api/upload")
async def upload_payment_screenshots(
    files: list[UploadFile] = File(...)
):

    allowed = {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp"
    }

    results = []

    approved = 0
    needs_review = 0
    pending_ai_count = 0
    duplicates = 0
    failed = 0

    for uploaded_file in files:

        original_name = (
            uploaded_file.filename
            or "payment"
        )

        safe_name = Path(
            original_name
        ).name

        extension = Path(
            safe_name
        ).suffix.lower()

        # ----------------------------------------------------
        # VALIDATE
        # ----------------------------------------------------

        if extension not in allowed:

            results.append({

                "filename":
                    safe_name,

                "status":
                    "REJECTED",

                "error":
                    "Unsupported image format."
            })

            failed += 1

            continue

        # ----------------------------------------------------
        # UNIQUE FILE NAME
        # ----------------------------------------------------

        destination = (
            UPLOAD_DIR /
            safe_name
        )

        counter = 1

        while destination.exists():

            destination = (
                UPLOAD_DIR /
                (
                    f"{Path(safe_name).stem}"
                    f"_{counter}"
                    f"{extension}"
                )
            )

            counter += 1

        # ----------------------------------------------------
        # SAVE SCREENSHOT
        # ----------------------------------------------------

        try:

            content = (
                await uploaded_file.read()
            )

            if not content:

                raise ValueError(
                    "Uploaded file is empty."
                )

            destination.write_bytes(
                content
            )

        except Exception as e:

            results.append({

                "filename":
                    safe_name,

                "status":
                    "FAILED",

                "error":
                    (
                        "Could not store screenshot: "
                        + str(e)
                    )
            })

            failed += 1

            continue

        # ----------------------------------------------------
        # PROCESS IMMEDIATELY
        # ----------------------------------------------------

        result = (
            process_screenshot_file(
                destination
            )
        )

        results.append(
            result
        )

        status = clean_text(
            result.get("status")
        ).upper()

        if status == "APPROVED":

            approved += 1

        elif status == "NEEDS_REVIEW":

            needs_review += 1

        elif status == "PENDING_AI":

            pending_ai_count += 1

        elif status == "DUPLICATE":

            duplicates += 1

        else:

            failed += 1

    return {

        "success": True,

        "uploaded":
            len(results),

        "approved":
            approved,

        "needs_review":
            needs_review,

        "pending_ai":
            pending_ai_count,

        "duplicates":
            duplicates,

        "failed":
            failed,

        "results":
            results
    }


# ============================================================
# ACCOUNTING AGENT
# ============================================================

@app.post("/api/ask")
async def ask_agent(
    payload: dict
):

    question = clean_text(
        payload.get(
            "question"
        )
    )

    if not question:

        return JSONResponse(
            status_code=400,
            content={
                "success": False,
                "error":
                    "Question is required."
            }
        )

    q = question.lower()

    summary_data = (
        accounting_summary()
    )

    payments = (
        get_approved_payments()
    )

    answer = None

    # ========================================================
    # NO APPROVED PAYMENTS
    # ========================================================

    if not payments:

        if summary_data["needs_review"] > 0:

            answer = (
                "There are currently no approved payments. "
                f"{summary_data['needs_review']} payment(s) "
                "are waiting for manual review, so they are "
                "not included in expenditure totals yet."
            )

        elif summary_data["pending_ai"] > 0:

            answer = (
                "There are currently no approved payments. "
                f"{summary_data['pending_ai']} screenshot(s) "
                "are still waiting for AI processing."
            )

        else:

            answer = (
                "There are currently no approved payments "
                "in the GCGW accounting workbook."
            )

    # ========================================================
    # TOTAL EXPENDITURE
    # ========================================================

    if answer is None and (
        "total expenditure" in q
        or "total expenditures" in q
        or "total expense" in q
        or "total expenses" in q
        or "total spent" in q
        or "total spending" in q
        or "overall expenditure" in q
        or "overall expense" in q
        or (
            "how much" in q
            and "spent" in q
            and "month" not in q
            and "today" not in q
        )
        or (
            "what are my" in q
            and "expenditure" in q
        )
        or (
            "what is my" in q
            and "expenditure" in q
        )
    ):

        answer = (
            "Total approved expenditure is "
            f"₹{summary_data['total_expenditure']:,.2f}."
        )

    # ========================================================
    # THIS MONTH
    # ========================================================

    if answer is None and (
        "this month" in q
        or "monthly expenditure" in q
        or "month expenditure" in q
        or "monthly expense" in q
    ):

        answer = (
            "Approved expenditure this month is "
            f"₹{summary_data['this_month']:,.2f}."
        )

    # ========================================================
    # TODAY
    # ========================================================

    if answer is None and (
        "today" in q
        and (
            "spend" in q
            or "spent" in q
            or "expenditure" in q
            or "expense" in q
            or "paid" in q
            or "payment" in q
        )
    ):

        answer = (
            "Approved expenditure today is "
            f"₹{summary_data['today']:,.2f}."
        )

    # ========================================================
    # NUMBER OF PAYMENTS
    # ========================================================

    if answer is None and (
        "how many payments" in q
        or "number of payments" in q
        or "approved payments" in q
    ):

        answer = (
            "There are "
            f"{summary_data['approved_payments']} "
            "approved payment(s)."
        )

    # ========================================================
    # NEEDS REVIEW COUNT
    # ========================================================

    if answer is None and (
        "needs review" in q
        or "need review" in q
        or "review payments" in q
    ):

        answer = (
            f"{summary_data['needs_review']} "
            "payment(s) currently need manual review."
        )

    # ========================================================
    # PENDING AI COUNT
    # ========================================================

    if answer is None and (
        "pending ai" in q
        or "pending payments" in q
        or "pending screenshot" in q
    ):

        answer = (
            f"{summary_data['pending_ai']} "
            "screenshot(s) are currently pending AI processing."
        )

    # ========================================================
    # HIGHEST SINGLE EXPENSE
    # ========================================================

    if answer is None and (
        "highest expense" in q
        or "largest expense" in q
        or "biggest expense" in q
        or "highest payment" in q
        or "largest payment" in q
    ):

        if payments:

            highest = max(
                payments,
                key=lambda x: safe_amount(
                    x.get("amount")
                )
            )

            answer = (
                "The highest approved payment is "
                f"₹{highest['amount']:,.2f}"
            )

            if highest.get(
                "paid_to"
            ):

                answer += (
                    f" paid to {highest['paid_to']}"
                )

            if highest.get(
                "project"
            ):

                answer += (
                    f" for {highest['project']}"
                )

            answer += "."

        else:

            answer = (
                "There are no approved payments yet."
            )

    # ========================================================
    # HIGHEST PROJECT
    # ========================================================

    if answer is None and (
        "highest" in q
        and "project" in q
    ):

        totals = (
            summary_data[
                "project_totals"
            ]
        )

        if totals:

            project = max(
                totals,
                key=totals.get
            )

            answer = (
                f"{project} currently has "
                "the highest approved expenditure "
                f"at ₹{totals[project]:,.2f}."
            )

        else:

            answer = (
                "There are no approved "
                "project expenses yet."
            )

    # ========================================================
    # PROJECT SEARCH
    # ========================================================

    if answer is None:

        for project, amount in (
            summary_data[
                "project_totals"
            ].items()
        ):

            if (
                project
                and
                project.lower() in q
            ):

                answer = (
                    f"Approved expenditure for "
                    f"{project} is "
                    f"₹{amount:,.2f}."
                )

                break

    # ========================================================
    # CATEGORY SEARCH
    # ========================================================

    if answer is None:

        for category, amount in (
            summary_data[
                "category_totals"
            ].items()
        ):

            if (
                category
                and
                category.lower() in q
            ):

                answer = (
                    f"Approved expenditure for "
                    f"{category} is "
                    f"₹{amount:,.2f}."
                )

                break

    # ========================================================
    # PAYEE SEARCH
    # ========================================================

    if answer is None:

        matching = []

        for payment in payments:

            payee = clean_text(
                payment.get(
                    "paid_to"
                )
            )

            if (
                payee
                and
                payee.lower() in q
            ):

                matching.append(
                    payment
                )

        if matching:

            amount = sum(
                safe_amount(
                    item.get(
                        "amount"
                    )
                )
                for item in matching
            )

            payee_name = (
                matching[0][
                    "paid_to"
                ]
            )

            answer = (
                f"There are {len(matching)} "
                f"approved payment(s) to "
                f"{payee_name}, totaling "
                f"₹{amount:,.2f}."
            )

    # ========================================================
    # SAFE FALLBACK
    # ========================================================

    if answer is None:

        answer = (
            "I could not match that question to an "
            "accounting calculation yet. You can ask things "
            "like: total expenditure, today's expenditure, "
            "this month's expenditure, highest expense, "
            "project expenditure, category expenditure, "
            "or payments to a particular person."
        )

    return {

        "success":
            True,

        "question":
            question,

        "answer":
            answer,

        "source":
            "GCGW approved Excel records",

        "approved_payments":
            summary_data[
                "approved_payments"
            ],

        "needs_review":
            summary_data[
                "needs_review"
            ],

        "pending_ai":
            summary_data[
                "pending_ai"
            ]
    }


# ============================================================
# START SERVER
# ============================================================

if __name__ == "__main__":

    import uvicorn

    port = int(
        os.environ.get(
            "PORT",
            8000
        )
    )

    uvicorn.run(
        "app:app",
        host="0.0.0.0",
        port=port
    )
